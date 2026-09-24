"""Парсеры документов → список блоков с путём заголовков.

Зачёркнутый текст и удалённые правки (track changes) оборачиваются в ~~...~~,
чтобы признак «удаляется» доходил до поиска и до LLM.
"""
from __future__ import annotations

import csv
import re
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Block:
    path: tuple[str, ...]     # цепочка заголовков
    text: str
    kind: str = "text"        # text | table
    header: str | None = None # для таблиц — строка заголовка, повторяется при разбиении
    loc: str | None = None    # страница / лист / слайд


@dataclass
class Parsed:
    title: str
    blocks: list[Block]
    warning: str | None = None


class HeadingStack:
    def __init__(self):
        self.items: list[tuple[int, str]] = []

    def push(self, level: int, title: str):
        while self.items and self.items[-1][0] >= level:
            self.items.pop()
        self.items.append((level, title))

    @property
    def path(self) -> tuple[str, ...]:
        return tuple(t for _, t in self.items)


def _clean(s: str) -> str:
    s = re.sub(r"[\u200b\u00ad]", "", s.replace("\xa0", " "))  # нулевой пробел, мягкий перенос
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"~~\s*~~", "", s)
    return s.strip()


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


# ---------------------------------------------------------------- DOCX

def load_docx(path: Path) -> Parsed:
    from docx import Document
    from docx.oxml.ns import qn

    W_P, W_TBL, W_TR, W_TC, W_R = qn("w:p"), qn("w:tbl"), qn("w:tr"), qn("w:tc"), qn("w:r")
    W_T, W_DELT, W_TAB, W_BR, W_CR = qn("w:t"), qn("w:delText"), qn("w:tab"), qn("w:br"), qn("w:cr")
    W_DEL, W_MOVEFROM = qn("w:del"), qn("w:moveFrom")
    W_RPR, W_PPR, W_VAL = qn("w:rPr"), qn("w:pPr"), qn("w:val")

    doc = Document(str(path))

    def on(el) -> bool:
        if el is None:
            return False
        v = el.get(W_VAL)
        return v is None or v.lower() not in ("0", "false", "off", "none")

    # style_id → уровень заголовка (0 = Title), с учётом наследования стилей
    style_level: dict[str, int | None] = {}
    heading_re = re.compile(r"^(heading|заголовок)\s*(\d)", re.I)

    def level_of_style(style) -> int | None:
        seen = 0
        while style is not None and seen < 10:
            name = (style.name or "").strip()
            if name.lower() in ("title", "название", "заголовок"):
                return 0
            m = heading_re.match(name)
            if m:
                return int(m.group(2))
            style = style.base_style
            seen += 1
        return None

    for st in doc.styles:
        try:
            style_level[st.style_id] = level_of_style(st)
        except Exception:
            pass

    def para_info(p):
        """(level|None, is_list, all_bold)"""
        ppr = p.find(W_PPR)
        level, is_list = None, False
        if ppr is not None:
            ps = ppr.find(qn("w:pStyle"))
            if ps is not None:
                sid = ps.get(W_VAL)
                level = style_level.get(sid)
                if sid and "list" in sid.lower():
                    is_list = True
            ol = ppr.find(qn("w:outlineLvl"))
            if level is None and ol is not None and ol.get(W_VAL, "").isdigit() and int(ol.get(W_VAL)) < 9:
                level = int(ol.get(W_VAL)) + 1
            if ppr.find(qn("w:numPr")) is not None:
                is_list = True
        runs = [r for r in p.iter(W_R) if any(c.tag == W_T and (c.text or "").strip() for c in r)]
        all_bold = bool(runs) and all(
            (r.find(W_RPR) is not None and on(r.find(W_RPR).find(qn("w:b")))) for r in runs
        )
        return level, is_list, all_bold

    def para_text(p) -> str:
        segs: list[tuple[str, bool]] = []
        for r in p.iter(W_R):
            rpr = r.find(W_RPR)
            struck = rpr is not None and (on(rpr.find(qn("w:strike"))) or on(rpr.find(qn("w:dstrike"))))
            deleted = any(a.tag in (W_DEL, W_MOVEFROM) for a in r.iterancestors())
            buf = []
            for ch in r:
                if ch.tag in (W_T, W_DELT):
                    buf.append(ch.text or "")
                elif ch.tag == W_TAB:
                    buf.append(" ")
                elif ch.tag in (W_BR, W_CR):
                    buf.append("\n")
            if buf:
                segs.append(("".join(buf), struck or deleted))
        merged: list[list] = []
        for t, s in segs:
            if merged and merged[-1][1] == s:
                merged[-1][0] += t
            else:
                merged.append([t, s])
        out = []
        for t, s in merged:
            core = t.strip()
            if s and core:
                lead, trail = t[: len(t) - len(t.lstrip())], t[len(t.rstrip()):]
                out.append(f"{lead}~~{core}~~{trail}")
            else:
                out.append(t)
        return _clean("".join(out))

    body = doc.element.body
    has_headings = any((para_info(p)[0] or 0) > 0 for p in body.iter(W_P))

    title = path.stem
    stack = HeadingStack()
    blocks: list[Block] = []

    for el in body.iterchildren():
        if el.tag == W_P:
            text = para_text(el)
            if not text:
                continue
            level, is_list, all_bold = para_info(el)
            if level == 0:
                title = " ".join(text.split())
                continue
            if level is not None:
                stack.push(level, " ".join(text.split()))
                continue
            # документ без стилей заголовков: короткий полностью жирный абзац — псевдозаголовок
            if not has_headings and all_bold and not is_list and len(text) <= 100 and not text.endswith((".", ":", ";")):
                stack.push(9, text)
                continue
            blocks.append(Block(stack.path, f"- {text}" if is_list else text))
        elif el.tag == W_TBL:
            rows = []
            for tr in el.findall(W_TR):
                cells = [_clean(" ".join(para_text(p) for p in tc.iter(W_P))) for tc in tr.findall(W_TC)]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(Block(stack.path, "\n".join(rows), kind="table", header=rows[0]))
    return Parsed(title, blocks)


# ---------------------------------------------------------------- PDF

def load_pdf(path: Path) -> Parsed:
    import pymupdf

    doc = pymupdf.open(str(path))
    title = (doc.metadata or {}).get("title") or path.stem
    toc_by_page: dict[int, list[tuple[int, str]]] = {}
    for lvl, t, page in doc.get_toc(simple=True):
        toc_by_page.setdefault(page, []).append((lvl, _clean(t)))

    stack = HeadingStack()
    blocks: list[Block] = []
    chars = 0
    for i, page in enumerate(doc, start=1):
        for lvl, t in toc_by_page.get(i, []):
            stack.push(lvl, t)
        for b in page.get_text("blocks", sort=True):
            if b[6] != 0:  # не текст
                continue
            txt = re.sub(r"[-\u00ad]\s*\n(?=[a-zа-яё])", "", b[4])  # склеиваем переносы по слогам
            txt = _clean(txt.replace("\n", " "))
            if len(txt) < 2:
                continue
            chars += len(txt)
            blocks.append(Block(stack.path, txt, loc=f"стр. {i}"))
    warning = None
    if doc.page_count and chars / doc.page_count < 50:
        warning = "почти нет текстового слоя — вероятно, скан (нужен OCR)"
    return Parsed(title, blocks, warning)


# ---------------------------------------------------------------- Markdown / TXT

def load_markdown(path: Path) -> Parsed:
    text = _read_text(path)
    title = path.stem
    stack = HeadingStack()
    blocks: list[Block] = []
    para: list[str] = []
    table: list[str] = []
    in_fence = False
    seen_h1 = False

    def flush_para():
        if para:
            t = _clean("\n".join(para)) if not in_fence else "\n".join(para)
            if t:
                blocks.append(Block(stack.path, t))
            para.clear()

    def flush_table():
        if table:
            rows = [r for r in table if not re.match(r"^\|?\s*:?-{2,}", r)]
            if rows:
                blocks.append(Block(stack.path, "\n".join(rows), kind="table", header=rows[0]))
            table.clear()

    for line in text.splitlines():
        if line.strip().startswith(("```", "~~~")):
            flush_table()
            if in_fence:
                para.append(line)
                blocks.append(Block(stack.path, "\n".join(para)))
                para.clear()
                in_fence = False
            else:
                flush_para()
                in_fence = True
                para.append(line)
            continue
        if in_fence:
            para.append(line)
            continue
        m = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", line)
        if m:
            flush_para(); flush_table()
            level, h = len(m.group(1)), _clean(re.sub(r"[*_`]", "", m.group(2)))
            if level == 1 and not seen_h1:
                seen_h1 = True
                title = h  # первый H1 — название документа, в путь раздела не входит
                continue
            stack.push(level, h)
            continue
        if line.strip().startswith("|"):
            flush_para()
            table.append(line.strip())
            continue
        flush_table()
        if not line.strip():
            flush_para()
        elif re.match(r"^\s*([-*+]|\d+[.)])\s+", line) and para:
            flush_para()
            para.append(line)
        else:
            para.append(line)
    flush_para(); flush_table()
    return Parsed(title, blocks)


def load_txt(path: Path, text: str | None = None) -> Parsed:
    text = _read_text(path) if text is None else text
    blocks = [Block((), _clean(p)) for p in re.split(r"\n\s*\n", text) if _clean(p)]
    return Parsed(path.stem, blocks)


def load_textutil(path: Path) -> Parsed:
    """.doc/.rtf/.odt — через встроенный в macOS textutil."""
    if sys.platform != "darwin":
        raise RuntimeError(f"{path.suffix} читается только в macOS (утилитой textutil); "
                           "пересохраните файл в .docx")
    out = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        capture_output=True, check=True,
    ).stdout.decode("utf-8", errors="replace")
    return load_txt(path, out)


# ---------------------------------------------------------------- HTML (в т.ч. экспорт Confluence)

def load_html(path: Path) -> Parsed:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_read_text(path), "lxml")
    for t in soup(["script", "style", "nav", "header", "footer", "noscript"]):
        t.decompose()
    for t in soup.find_all(["s", "del", "strike"]):
        inner = t.get_text(" ", strip=True)
        t.replace_with(f"~~{inner}~~" if inner else "")
    title = _clean(soup.title.get_text()) if soup.title and soup.title.get_text(strip=True) else path.stem

    stack = HeadingStack()
    blocks: list[Block] = []
    leaf = {"p", "li", "pre", "dt", "dd", "blockquote"}
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "table", *leaf]):
        if el.find_parent(["table", "li", "pre", "p", "blockquote"]):
            continue
        name = el.name
        if name[0] == "h" and name[1:].isdigit():
            h = _clean(el.get_text(" "))
            if h:
                stack.push(int(name[1:]), h)
            continue
        if name == "table":
            rows = []
            for tr in el.find_all("tr"):
                cells = [_clean(c.get_text(" ")) for c in tr.find_all(["td", "th"])]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(Block(stack.path, "\n".join(rows), kind="table", header=rows[0]))
            continue
        txt = el.get_text("\n" if name == "pre" else " ")
        txt = txt.strip() if name == "pre" else _clean(txt)
        if txt:
            blocks.append(Block(stack.path, f"- {txt}" if name == "li" else txt))
    return Parsed(title, blocks)


# ---------------------------------------------------------------- Таблицы / презентации

def _rows_to_blocks(rows: list[list[str]], sheet: str | None, loc: str | None) -> list[Block]:
    lines = []
    for r in rows:
        cells = [_clean(str(c)) if c is not None else "" for c in r]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            lines.append(" | ".join(cells))
    if not lines:
        return []
    return [Block((sheet,) if sheet else (), "\n".join(lines), kind="table", header=lines[0], loc=loc)]


def load_xlsx(path: Path) -> Parsed:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    for ws in wb.worksheets:
        blocks += _rows_to_blocks(list(ws.iter_rows(values_only=True)), ws.title, f"лист «{ws.title}»")
    wb.close()
    return Parsed(path.stem, blocks)


def load_csv(path: Path) -> Parsed:
    text = _read_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return Parsed(path.stem, _rows_to_blocks(list(csv.reader(text.splitlines(), dialect)), None, None))


def load_pptx(path: Path) -> Parsed:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[Block] = []
    for i, slide in enumerate(prs.slides, start=1):
        stitle = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            stitle = _clean(slide.shapes.title.text_frame.text)
        spath = (f"Слайд {i}" + (f": {stitle}" if stitle else ""),)
        loc = f"слайд {i}"
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                t = _clean("\n".join(p.text for p in shape.text_frame.paragraphs))
                if t:
                    blocks.append(Block(spath, t, loc=loc))
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                blocks += [Block(spath, b.text, "table", b.header, loc) for b in _rows_to_blocks(rows, None, None)]
        if slide.has_notes_slide:
            n = _clean(slide.notes_slide.notes_text_frame.text)
            if n:
                blocks.append(Block(spath, f"Заметки: {n}", loc=loc))
    return Parsed(path.stem, blocks)


LOADERS = {
    ".docx": load_docx,
    ".pdf": load_pdf,
    ".md": load_markdown, ".markdown": load_markdown,
    ".txt": load_txt,
    ".html": load_html, ".htm": load_html,
    ".xlsx": load_xlsx, ".xlsm": load_xlsx,
    ".csv": load_csv,
    ".pptx": load_pptx,
    ".doc": load_textutil, ".rtf": load_textutil, ".odt": load_textutil,
}


def load(path: Path) -> Parsed:
    return LOADERS[path.suffix.lower()](path)
